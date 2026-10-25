"""Synthetic-only operational contract; PostgreSQL race cases run in full CI."""
import copy
import io
import os
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from threading import Barrier
from unittest.mock import patch
from zipfile import ZIP_DEFLATED, ZipFile

from django.conf import settings
from django.contrib import admin
from django.contrib.admin.models import LogEntry
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.core.exceptions import PermissionDenied, ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import close_old_connections
from django.test import Client, TestCase, TransactionTestCase, override_settings, skipUnlessDBFeature
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from .handovers import claim_handover, issue_pilot_invitation
from .inventory_import import (ImportProblem, MAX_BYTES, PERMISSIONS, apply_preview,
                               make_preview, parse_upload, template_bytes)
from .models import (BotanicalSpecies, CareGuideVersion, HandoverInvitation,
                     InstallationBatch, InstallationManifestRow, InventoryImportReceipt,
                     PlantEvent, PlantUnit, Space, SpaceMember)


class ImportFixture:
    def setUp(self):
        self.staff = get_user_model().objects.create_superuser(username="qa_import_staff", email="import@example.test", password="qa-only")
        self.space = Space.objects.create(owner=self.staff, name="FICTIONAL QA installation")
        self.batch = InstallationBatch.objects.create(space=self.space, prepared_by=self.staff, reference_code="QA-IMPORT-001")
        self.client.force_login(self.staff)
        self.url = reverse("admin:plants_inventory_import")

    def upload(self, data=None, name="qa.xlsx"):
        return SimpleUploadedFile(name, data if data is not None else template_bytes(self.batch))

    def document(self, verified=False):
        doc = parse_upload(self.upload(), self.batch)
        if verified:
            doc["PlantUnits"][0].update(verified_species_key="qa-fictional-species", identity_reference="QA-only independent identity inspection record 01")
        return doc

    def approve(self, doc=None):
        plan, token = make_preview(self.staff, self.batch, doc or self.document())
        self.assertFalse(plan["errors"], plan)
        return apply_preview(self.staff, token, confirmed=True, identities_confirmed=True)


@override_settings(STORAGES={**settings.STORAGES, "staticfiles": {"BACKEND": "django.contrib.staticfiles.storage.StaticFilesStorage"}})
class InventoryImportTests(ImportFixture, TestCase):
    def test_admin_preview_no_writes_then_explicit_atomic_approval_and_replay(self):
        counts = lambda: (PlantUnit.objects.count(), BotanicalSpecies.objects.count(), CareGuideVersion.objects.count(), InventoryImportReceipt.objects.count(), LogEntry.objects.count())
        initial = counts()
        response = self.client.post(self.url, {"batch": self.batch.pk, "file": self.upload()})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(counts(), initial)
        self.assertContains(response, str(self.space.pk))
        self.assertContains(response, "Unreviewed draft")
        self.assertContains(response, "FICTIONAL QA")
        token = response.context["approval"].initial["preview"]
        rejected = self.client.post(self.url, {"preview": token})
        self.assertContains(rejected, "This field is required")
        self.assertEqual(counts(), initial)
        applied = self.client.post(self.url, {"preview": token, "confirmed": "on"})
        self.assertContains(applied, "Import committed atomically")
        plant = PlantUnit.objects.get()
        original_id = plant.pk
        guide = CareGuideVersion.objects.get()
        self.assertFalse(guide.published)
        self.assertIsNone(guide.reviewer)
        self.assertIsNone(guide.reviewed_at)
        self.assertIsNone(plant.verified_species)
        self.assertEqual(InstallationManifestRow.objects.get().plant_id, original_id)
        self.assertFalse(SpaceMember.objects.exists())
        self.assertFalse(HandoverInvitation.objects.exists())
        self.assertFalse(PlantEvent.objects.exists())
        with patch("plants.inventory_import.signing.dumps", return_value="x" * (MAX_BYTES + 1)):
            with self.assertRaisesMessage(ImportProblem, "Split this workbook"):
                make_preview(self.staff, self.batch, self.document())
        replay = self.client.post(self.url, {"preview": token, "confirmed": "on"})
        self.assertContains(replay, "already applied")
        self.assertEqual(PlantUnit.objects.get().pk, original_id)
        self.assertEqual(InventoryImportReceipt.objects.count(), 1)
        _, replay_token = make_preview(self.staff, self.batch, self.document())
        self.assertFalse(apply_preview(self.staff, replay_token, confirmed=True)[1])

    def test_verified_identity_requires_proof_and_actual_staff_confirmation(self):
        document = self.document(verified=True)
        _, token = make_preview(self.staff, self.batch, document)
        with self.assertRaises(ImportProblem):
            apply_preview(self.staff, token, confirmed=True)
        self.assertFalse(PlantUnit.objects.exists())
        self.approve(document)
        plant = PlantUnit.objects.get()
        self.assertEqual(plant.identity_verified_by, self.staff)
        self.assertTrue(plant.identity_verified_at)
        self.assertEqual(plant.identity_reference, document["PlantUnits"][0]["identity_reference"])
        self.assertEqual(plant.verified_species.key, "qa-fictional-species")

    def test_plant_update_uses_batch_source_key_preserves_uuid_photo_history_and_omitted_rows(self):
        self.approve(self.document(verified=True))
        plant = PlantUnit.objects.get()
        plant.cover_photo = "plantupia/qa-existing.png"
        plant.save()
        extra = PlantUnit.objects.create(space=self.space, label="Omitted QA plant")
        original_id = plant.pk
        doc = self.document()
        doc["PlantUnits"][0]["label"] = "QA اسم النبات — edited"
        plan, token = make_preview(self.staff, self.batch, doc)
        self.assertIn("update", plan["summary"])
        apply_preview(self.staff, token, confirmed=True)
        plant.refresh_from_db()
        self.assertEqual(plant.pk, original_id)
        self.assertEqual(plant.cover_photo.name, "plantupia/qa-existing.png")
        self.assertEqual(plant.label, "QA اسم النبات — edited")
        self.assertIsNone(plant.verified_species_id)
        self.assertIsNone(plant.identity_verified_by)
        self.assertEqual(plant.identity_reference, "")
        self.assertTrue(PlantUnit.objects.filter(pk=extra.pk).exists())

    def test_stale_preview_never_overwrites_intervening_inventory_change(self):
        self.approve()
        doc = self.document()
        doc["PlantUnits"][0]["label"] = "Incoming edit"
        _, token = make_preview(self.staff, self.batch, doc)
        PlantUnit.objects.update(label="Intervening staff edit")
        with self.assertRaisesMessage(ImportProblem, "changed since preview"):
            apply_preview(self.staff, token, confirmed=True)
        self.assertEqual(PlantUnit.objects.get().label, "Intervening staff edit")
        self.assertEqual(InventoryImportReceipt.objects.count(), 1)

    def test_late_failure_rolls_back_species_guides_plants_manifest_and_receipt(self):
        _, token = make_preview(self.staff, self.batch, self.document())
        with patch.object(InventoryImportReceipt.objects, "create", side_effect=ValidationError("QA injected commit failure")):
            with self.assertRaises(ImportProblem):
                apply_preview(self.staff, token, confirmed=True)
        for model in (BotanicalSpecies, CareGuideVersion, PlantUnit, InstallationManifestRow, InventoryImportReceipt):
            self.assertEqual(model.objects.count(), 0, model)

    def test_new_guide_versions_only_existing_drafts_and_published_content_never_overwritten(self):
        self.approve()
        document = self.document()
        document["CareGuides"][0]["instructions"] = "Changed QA advice"
        for published in (False, True):
            guide = CareGuideVersion.objects.get()
            if published:
                guide.reviewer, guide.reviewed_at = self.staff, timezone.now()
                guide.published = True
                guide.save()
            plan, token = make_preview(self.staff, self.batch, document)
            self.assertTrue(plan["errors"])
            self.assertIsNone(token)
        document["CareGuides"][0]["version"] = 2
        self.approve(document)
        self.assertFalse(CareGuideVersion.objects.get(version=2).published)

    def test_permission_scope_signed_actor_tamper_expiry_and_csrf(self):
        _, token = make_preview(self.staff, self.batch, self.document())
        other = get_user_model().objects.create_user(username="qa_other_importer", email="other@example.test", is_staff=True)
        self.client.force_login(other)
        self.assertEqual(self.client.get(self.url).status_code, 403)
        permissions = [p.split(".")[1] for p in PERMISSIONS]
        other.user_permissions.set(Permission.objects.filter(codename__in=permissions, content_type__app_label="plants"))
        other = get_user_model().objects.get(pk=other.pk)
        self.client.force_login(other)
        page = self.client.get(self.url)
        self.assertNotContains(page, self.space.name)
        self.assertEqual(self.client.get(self.url, {"batch": self.batch.pk}).status_code, 404)
        self.assertEqual(self.client.post(self.url, {"batch": self.batch.pk, "file": self.upload()}).context["form"].errors.keys(), {"batch"})
        with self.assertRaises(PermissionDenied):
            apply_preview(other, token, confirmed=True)
        with self.assertRaises(ImportProblem):
            apply_preview(self.staff, token + "tampered", confirmed=True)
        document = self.document()
        with patch("django.core.signing.TimestampSigner.timestamp", return_value="1"):
            _, expired = make_preview(self.staff, self.batch, document)
        with self.assertRaises(ImportProblem):
            apply_preview(self.staff, expired, confirmed=True)
        csrf = Client(enforce_csrf_checks=True)
        csrf.force_login(self.staff)
        self.assertEqual(csrf.post(self.url, {"batch": self.batch.pk, "file": self.upload()}).status_code, 403)
        csrf.get(self.url)
        result = csrf.post(self.url, {"batch": self.batch.pk, "file": self.upload(), "csrfmiddlewaretoken": csrf.cookies["csrftoken"].value})
        self.assertIsNotNone(result.context["plan"])

    def test_nonstaff_anonymous_and_missing_granular_permission_cannot_import(self):
        self.client.logout()
        self.assertEqual(self.client.get(self.url).status_code, 302)
        self.staff.is_superuser = False
        self.staff.save()
        self.staff.user_permissions.set(Permission.objects.filter(content_type__app_label="plants").exclude(codename="add_careguideversion"))
        self.client.force_login(self.staff)
        self.assertEqual(self.client.get(self.url).status_code, 403)
        self.staff.is_staff = False
        self.staff.save()
        self.client.force_login(self.staff)
        self.assertEqual(self.client.get(self.url).status_code, 302)

    def test_reviewed_claimed_or_moved_destination_and_cross_space_manifest_refused(self):
        document = self.document()
        _, token = make_preview(self.staff, self.batch, document)
        for status in ("ready", "claimed"):
            InstallationBatch.objects.filter(pk=self.batch.pk).update(status=status)
            with self.assertRaises(PermissionDenied):
                apply_preview(self.staff, token, confirmed=True)
        InstallationBatch.objects.filter(pk=self.batch.pk).update(status="draft")
        elsewhere = Space.objects.create(owner=self.staff, name="Other QA space")
        plant = PlantUnit.objects.create(space=elsewhere, label="QA outside destination")
        InstallationManifestRow.objects.create(batch=self.batch, plant=plant, source_key="qa-other")
        plan, token = make_preview(self.staff, self.batch, document)
        self.assertTrue(plan["errors"])
        self.assertIsNone(token)

    def test_csv_is_inventory_only_and_destination_bound(self):
        document = parse_upload(self.upload(template_bytes(self.batch, True), "qa.csv"), self.batch)
        self.assertEqual(document["Species"], [])
        self.assertEqual(document["CareGuides"], [])
        self.approve(document)
        self.assertEqual(PlantUnit.objects.count(), 1)
        self.assertFalse(CareGuideVersion.objects.exists())
        wrong = template_bytes(self.batch, True).replace(str(self.space.pk).encode(), b"incorrect-space")
        with self.assertRaises(ImportProblem):
            parse_upload(self.upload(wrong, "qa.csv"), self.batch)

    def test_publication_records_actual_reviewer_withdrawal_remains_immutable_and_api_matches_android(self):
        self.approve(self.document(verified=True))
        guide = CareGuideVersion.objects.get()
        model_admin = admin.site._registry[CareGuideVersion]
        self.assertIn("published", model_admin.get_readonly_fields(None, guide))
        url = reverse("admin:plants_careguideversion_changelist")
        response = self.client.post(url, {"action": "review_and_publish", "_selected_action": [guide.pk]})
        self.assertEqual(response.status_code, 302)
        guide.refresh_from_db()
        self.assertTrue(guide.published)
        self.assertEqual(guide.reviewer, self.staff)
        self.assertTrue(guide.reviewed_at)
        api = APIClient()
        api.force_authenticate(self.staff)
        plant = PlantUnit.objects.get()
        body = api.get(reverse("plantupia-care-guide", args=[plant.pk])).json()
        self.assertEqual(body["state"], "available")
        self.assertEqual(body["placement"], plant.placement_note)
        self.assertEqual(body["guide"]["version"], 1)
        self.assertEqual(body["guide"]["reviewer"], self.staff.username)
        self.assertEqual(body["guide"]["source_reference"], guide.source_reference)
        self.client.post(url, {"action": "withdraw", "_selected_action": [guide.pk]})
        guide.refresh_from_db()
        guide.instructions = "Changed after withdrawal"
        with self.assertRaises(ValidationError):
            guide.save()
        self.assertEqual(api.get(reverse("plantupia-care-guide", args=[plant.pk])).json()["state"], "unverified")

    def test_import_review_claim_uses_same_uuid_and_private_guide(self):
        self.approve(self.document(verified=True))
        guide = CareGuideVersion.objects.get()
        self.client.post(reverse("admin:plants_careguideversion_changelist"), {"action": "review_and_publish", "_selected_action": [guide.pk]})
        self.client.post(reverse("admin:plants_installationbatch_changelist"), {"action": "mark_inventory_ready", "_selected_action": [self.batch.pk]})
        self.batch.refresh_from_db()
        self.assertEqual(self.batch.status, "ready")
        customer = get_user_model().objects.create_user(username="qa_import_customer", email="customer@example.test")
        plant_id = PlantUnit.objects.get().pk
        with patch.dict(os.environ, {"PLANTUPIA_HANDOVER_QA_ENABLED": "1"}):
            code, _ = issue_pilot_invitation(batch_id=self.batch.pk, recipient=customer, staff=self.staff, mode="customer_owned")
            claim_handover(user=customer, token=code)
        self.assertEqual(PlantUnit.objects.get().pk, plant_id)
        api = APIClient()
        api.force_authenticate(customer)
        self.assertEqual(api.get(reverse("plantupia-care-guide", args=[plant_id])).json()["state"], "available")
        api.force_authenticate(self.staff)
        self.assertEqual(api.get(reverse("plantupia-care-guide", args=[plant_id])).status_code, 404)

    def test_template_download_discovery_and_html_escape(self):
        for name in ("plants_installationbatch_changelist", "plants_plantunit_changelist", "plants_careguideversion_changelist"):
            self.assertContains(self.client.get(reverse("admin:" + name)), "Import inventory")
        url = reverse("admin:plants_inventory_template", args=[self.batch.pk])
        template = self.client.get(url)
        self.assertIn("no-store", template["Cache-Control"])
        self.assertEqual(template["Content-Disposition"], 'attachment; filename="plantupia-inventory-v1.xlsx"')
        book = load_workbook(io.BytesIO(template.content))
        book["PlantUnits"]["B2"] = '<script>alert("qa")</script> نبات'
        out = io.BytesIO()
        book.save(out)
        response = self.client.post(self.url, {"batch": self.batch.pk, "file": self.upload(out.getvalue(), "../../qa.xlsx")})
        self.assertNotContains(response, '<script>alert("qa")</script>')
        self.assertContains(response, "&lt;script&gt;")


class InventoryParserTests(ImportFixture, TestCase):
    def mutate(self, callback):
        book = load_workbook(io.BytesIO(template_bytes(self.batch)))
        callback(book)
        out = io.BytesIO()
        book.save(out)
        return self.upload(out.getvalue())

    def test_rejects_bad_columns_keys_sources_formula_types_dates_and_rows(self):
        changes = [
            lambda b: setattr(b["PlantUnits"]["H1"], "value", "plant_id"),
            lambda b: b["PlantUnits"].append(["qa-unapproved-001", "Duplicate QA", "individual"]),
            lambda b: b["Species"].append(["qa-fictional-species", "Duplicate QA identity"]),
            lambda b: b["CareGuides"].append([c.value for c in b["CareGuides"][2]]),
            lambda b: setattr(b["Metadata"]["B4"], "value", "wrong-space"),
            lambda b: setattr(b["CareGuides"]["C2"], "value", "http://example.test/untrusted"),
            lambda b: setattr(b["CareGuides"]["B2"], "value", 0),
            lambda b: setattr(b["PlantUnits"]["B2"], "value", "=WEBSERVICE(\"https://example.test\")"),
            lambda b: setattr(b["PlantUnits"]["B2"], "value", "@formula"),
            lambda b: setattr(b["PlantUnits"]["B2"], "value", datetime(2026, 1, 1)),
            lambda b: setattr(b["PlantUnits"]["B2"], "value", True),
            lambda b: setattr(b["PlantUnits"]["A2"], "value", 123),
            lambda b: setattr(b["PlantUnits"]["B2"], "value", ""),
            lambda b: setattr(b["PlantUnits"]["B2"], "value", "x" * 121),
            lambda b: setattr(b["PlantUnits"]["E2"], "value", "qa-fictional-species"),
            lambda b: b.create_sheet("Approved"),
            lambda b: setattr(b["PlantUnits"]["A502"], "value", "too-many-rows"),
        ]
        for change in changes:
            with self.subTest(change=change):
                with self.assertRaises(ImportProblem):
                    parse_upload(self.mutate(change), self.batch)
        self.assertFalse(PlantUnit.objects.exists())

    def test_corrupt_oversized_legacy_non_utf8_csv_and_external_relationships_rejected(self):
        for name, data in (("qa.xlsx", b"not-a-zip"), ("qa.xls", b"legacy"), ("qa.xlsm", b"macro"), ("qa.csv", b"\xff"), ("qa.csv", b"x" * (MAX_BYTES + 1))):
            with self.subTest(name=name, size=len(data)), self.assertRaises(ImportProblem):
                parse_upload(self.upload(data, name), self.batch)
        original = template_bytes(self.batch)
        for name, content in (("xl/externalLinks/externalLink1.xml", b"<x/>"),
                              ("xl/vbaProject.bin", b"qa"),
                              ("xl/media/image1.png", b"qa-embedded-image"),
                              ("custom/external.rels", b'<Relationships><Relationship TargetMode="External" Target="file:///qa"/></Relationships>'),
                              ("custom/entity.xml", b'<!DOCTYPE x [<!ENTITY a "qa">]><x>&a;</x>')):
            out = io.BytesIO()
            with ZipFile(io.BytesIO(original)) as source, ZipFile(out, "w", ZIP_DEFLATED) as dest:
                for entry in source.infolist():
                    dest.writestr(entry, source.read(entry))
                dest.writestr(name, content)
            with self.subTest(name=name), self.assertRaises(ImportProblem):
                parse_upload(self.upload(out.getvalue()), self.batch)


class InventoryConcurrencyTests(ImportFixture, TransactionTestCase):
    @skipUnlessDBFeature("has_select_for_update")
    def test_stale_draft_editor_cannot_overwrite_concurrently_published_content(self):
        self.approve()
        guide_id = CareGuideVersion.objects.get().pk
        barrier = Barrier(2)

        def save(publish):
            close_old_connections()
            try:
                guide = CareGuideVersion.objects.get(pk=guide_id)
                if publish:
                    guide.published, guide.reviewer, guide.reviewed_at = True, self.staff, timezone.now()
                else:
                    guide.instructions = "QA concurrent editor content"
                barrier.wait(timeout=10)
                try:
                    guide.save()
                    return "saved"
                except ValidationError:
                    return "frozen"
            finally:
                close_old_connections()

        with ThreadPoolExecutor(max_workers=2) as pool:
            published, edited = list(pool.map(save, [True, False]))
        self.assertEqual(published, "saved")
        self.assertIn(edited, ("saved", "frozen"))
        guide = CareGuideVersion.objects.get(pk=guide_id)
        self.assertTrue(guide.published)
        self.assertEqual(guide.instructions, self.document()["CareGuides"][0]["instructions"])

    @skipUnlessDBFeature("has_select_for_update")
    def test_two_simultaneous_approvals_same_preview_commit_one_receipt(self):
        _, token = make_preview(self.staff, self.batch, self.document(verified=True))
        barrier = Barrier(2)

        def apply(_):
            close_old_connections()
            try:
                user = get_user_model().objects.get(pk=self.staff.pk)
                barrier.wait(timeout=10)
                receipt, created = apply_preview(user, token, confirmed=True, identities_confirmed=True)
                return receipt.pk, created
            finally:
                close_old_connections()

        with ThreadPoolExecutor(max_workers=2) as pool:
            results = list(pool.map(apply, range(2)))
        self.assertEqual(results[0][0], results[1][0])
        self.assertEqual(sorted(r[1] for r in results), [False, True])
        for model in (BotanicalSpecies, CareGuideVersion, PlantUnit, InstallationManifestRow, InventoryImportReceipt):
            self.assertEqual(model.objects.count(), 1)

    @skipUnlessDBFeature("has_select_for_update")
    def test_two_different_previews_same_destination_one_wins_other_requires_fresh_preview(self):
        first = self.document()
        second = copy.deepcopy(first)
        second["PlantUnits"][0]["label"] = "Other incoming QA label"
        tokens = [make_preview(self.staff, self.batch, doc)[1] for doc in (first, second)]
        barrier = Barrier(2)

        def apply(token):
            close_old_connections()
            try:
                user = get_user_model().objects.get(pk=self.staff.pk)
                barrier.wait(timeout=10)
                try:
                    apply_preview(user, token, confirmed=True)
                    return "committed"
                except ImportProblem:
                    return "stale"
            finally:
                close_old_connections()

        with ThreadPoolExecutor(max_workers=2) as pool:
            self.assertEqual(sorted(pool.map(apply, tokens)), ["committed", "stale"])
        self.assertEqual(PlantUnit.objects.count(), 1)
        self.assertEqual(InventoryImportReceipt.objects.count(), 1)

    @skipUnlessDBFeature("has_select_for_update")
    def test_import_racing_ready_review_never_mutates_already_reviewed_inventory(self):
        self.approve()
        document = self.document()
        document["PlantUnits"][0]["label"] = "QA amendment before review"
        _, token = make_preview(self.staff, self.batch, document)
        barrier = Barrier(2)

        def operation(review):
            close_old_connections()
            try:
                user = get_user_model().objects.get(pk=self.staff.pk)
                if review:
                    client = Client()
                    client.force_login(user)
                    barrier.wait(timeout=10)
                    return client.post(reverse("admin:plants_installationbatch_changelist"), {
                        "action": "mark_inventory_ready", "_selected_action": [self.batch.pk]}).status_code
                barrier.wait(timeout=10)
                try:
                    apply_preview(user, token, confirmed=True)
                    return "imported"
                except (ImportProblem, PermissionDenied):
                    return "locked"
            finally:
                close_old_connections()

        with ThreadPoolExecutor(max_workers=2) as pool:
            imported, reviewed = list(pool.map(operation, [False, True]))
        self.assertEqual(reviewed, 302)
        self.batch.refresh_from_db()
        self.assertEqual(self.batch.status, "ready")
        self.assertEqual(PlantUnit.objects.get().label,
                         "QA amendment before review" if imported == "imported" else "FICTIONAL QA — NOT APPROVED")

    @skipUnlessDBFeature("has_select_for_update")
    def test_different_batches_racing_new_shared_taxonomy_roll_back_loser_then_retry_safely(self):
        other_space = Space.objects.create(owner=self.staff, name="Second FICTIONAL QA space")
        other = InstallationBatch.objects.create(space=other_space, prepared_by=self.staff, reference_code="QA-IMPORT-002")
        document = self.document()
        tokens = [make_preview(self.staff, batch, document)[1] for batch in (self.batch, other)]
        barrier = Barrier(2)

        def apply(token):
            close_old_connections()
            try:
                user = get_user_model().objects.get(pk=self.staff.pk)
                barrier.wait(timeout=10)
                try:
                    apply_preview(user, token, confirmed=True)
                    return "committed"
                except ImportProblem:
                    return "conflict"
            finally:
                close_old_connections()

        with ThreadPoolExecutor(max_workers=2) as pool:
            self.assertEqual(sorted(pool.map(apply, tokens)), ["committed", "conflict"])
        self.assertEqual(InventoryImportReceipt.objects.count(), 1)
        for batch in (self.batch, other):
            _, token = make_preview(self.staff, batch, document)
            apply_preview(self.staff, token, confirmed=True)
        self.assertEqual(PlantUnit.objects.count(), 2)
        self.assertEqual(InstallationManifestRow.objects.count(), 2)
        self.assertEqual(BotanicalSpecies.objects.count(), 1)
        self.assertEqual(CareGuideVersion.objects.count(), 1)
        self.assertEqual(InventoryImportReceipt.objects.count(), 2)
