"""Fresh local CEO rehearsal, composed with the existing guarded company fixture."""
from io import BytesIO
from pathlib import Path

from django.contrib.auth import get_user_model
from django.core.management import call_command
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from plants import demo_services
from plants.inventory_import import template_bytes
from plants.models import CatalogProduct, DemoOffer, InstallationBatch, Space
from plants.qa_demo_care import SAMPLE_CARE_GUIDES


class Command(BaseCommand):
    help = "Create ordinary staff/customer accounts, concept offers, fictional products and a destination-bound CEO workbook in EMPTY local SQLite QA only."

    def add_arguments(self, parser):
        parser.add_argument("--credentials-file", required=True)

    def handle(self, *args, **options):
        if not demo_services.enabled():
            raise CommandError("Explicit isolated CEO QA and company-dashboard flags are required.")
        credentials = Path(options["credentials_file"]).expanduser().resolve()
        workbook_path = credentials.parent / "fictional-ceo-inventory.xlsx"
        if credentials.exists() or workbook_path.exists():
            raise CommandError("Choose a fresh private directory; existing credentials/workbook are preserved.")
        wrote_credentials = False
        wrote_workbook = False
        try:
            with transaction.atomic():
                # This existing command refuses hosted/non-SQLite/nonempty databases,
                # credentials in Git and disabled handovers before creating anything.
                call_command("seed_local_plantupia_company_qa", credentials_file=str(credentials))
                wrote_credentials = True
                operator = get_user_model().objects.get(username="qa_company_operator")
                space = Space.objects.create(owner=operator, name="FICTIONAL QA · CEO delivery", description="Disposable CEO rehearsal. Not a real customer or installation.")
                batch = InstallationBatch.objects.create(space=space, prepared_by=operator, reference_code="QA-CEO-DELIVERY")
                book = load_workbook(BytesIO(template_bytes(batch)))
                for name in ("Species", "CareGuides", "PlantUnits"):
                    book[name].delete_rows(2, book[name].max_row)
                for number, sample in enumerate(SAMPLE_CARE_GUIDES, 1):
                    book["Species"].append([sample["key"], sample["scientific_name"]])
                    book["CareGuides"].append([sample["key"], 1, sample["source_url"], sample["source_reference"], sample["instructions"], sample["context_notes"]])
                    book["PlantUnits"].append([f"CEO-QA-{number:02}", f"FICTIONAL QA supplied plant {number:02}", "individual", sample["scientific_name"], sample["key"],
                        "SYNTHETIC identity assumption for software QA; NOT independently verified real botanical evidence.", f"Synthetic QA position {number:02}"])
                with workbook_path.open("xb") as output:
                    wrote_workbook = True
                    book.save(output)
                DemoOffer.objects.create(key="ai-care", name="AI Care", proposed_monthly_price="79.00", ai_daily_limit=20)
                DemoOffer.objects.create(key="expert-care", name="Expert Care", proposed_monthly_price="199.00", ai_daily_limit=35, human_requests_per_30_days=3)
                for name, kind, price in (("indoor plant", "plant", "125.50"), ("ceramic planter", "planter", "80.00"), ("inspection accessory", "accessory", "35.99")):
                    CatalogProduct.objects.create(name="FICTIONAL QA · " + name, kind=kind, proposed_price=price,
                        description="Fictional product for no-charge software QA. Not official Plantupia merchandise, live inventory, a delivery or a commercial offer.", qa_stock=10)
            self.stdout.write("CEO QA fixture created; nothing was imported, published, scheduled, entitled, ordered or handed over automatically.")
            self.stdout.write("Private workbook: " + str(workbook_path))
        except BaseException:
            if wrote_workbook:
                workbook_path.unlink(missing_ok=True)
            if wrote_credentials:
                credentials.unlink(missing_ok=True)
            raise
