"""Bounded, read-only parsing/planning and explicit transactional staff import.

No remote URLs are fetched, no filenames are stored and no customer payload is
logged. The signed preview lives only in the staff browser, expires in 30 minutes
and is bound to actor, destination, content and the observed database state.
"""
import csv
import hashlib
import io
import json
import re
from collections import Counter
from zipfile import ZipFile

from defusedxml.ElementTree import fromstring
from django.core import signing
from django.core.exceptions import PermissionDenied, ValidationError
from django.core.validators import URLValidator
from django.db import IntegrityError, transaction
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from .models import (BotanicalSpecies, CareGuideVersion, InstallationBatch,
                     InstallationManifestRow, InventoryImportReceipt, PlantUnit, Space)

VERSION = "plantupia-inventory-v1"
MAX_BYTES = 2 * 1024 * 1024
MAX_EXPANDED = 16 * 1024 * 1024
MAX_ROWS = 500
MAX_CELL = 8000
SALT = "plants.inventory.preview.v1"
SCHEMA = {
    "Metadata": ("field", "value"),
    "Species": ("key", "scientific_name"),
    "CareGuides": ("species_key", "version", "source_url", "source_reference", "instructions", "context_notes"),
    "PlantUnits": ("source_key", "label", "unit_type", "species_label", "verified_species_key", "identity_reference", "placement_note"),
}
PERMISSIONS = ("plants.import_inventory", "plants.change_installationbatch",
               "plants.view_space", "plants.add_plantunit", "plants.change_plantunit",
               "plants.add_installationmanifestrow", "plants.change_installationmanifestrow",
               "plants.add_botanicalspecies", "plants.view_botanicalspecies",
               "plants.add_careguideversion", "plants.view_careguideversion")


class ImportProblem(ValueError):
    pass


def eligible_batches(user):
    qs = InstallationBatch.objects.filter(status=InstallationBatch.Status.DRAFT, space__owner__is_staff=True)
    return qs if user.is_superuser else qs.filter(prepared_by=user, space__owner=user)


def require_permissions(user):
    if not user.is_active or not user.is_staff or not user.has_perms(PERMISSIONS):
        raise PermissionDenied


def authorize(user, batch):
    require_permissions(user)
    if not eligible_batches(user).filter(pk=batch.pk).exists():
        raise PermissionDenied
    if batch.status != InstallationBatch.Status.DRAFT or not batch.space.owner.is_staff:
        raise ImportProblem("Only an existing staff-owned DRAFT installation can receive an import.")


def digest(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True, ensure_ascii=False, default=str).encode()).hexdigest()


def text_cell(value, where):
    if value is None:
        return ""
    if isinstance(value, bool) or not isinstance(value, (str, int)):
        raise ImportProblem(f"{where}: use text cells (only guide version may be numeric); dates and other cell types are unsupported.")
    value = str(value).strip()
    if len(value) > MAX_CELL or any(ord(c) < 32 and c not in "\n\r\t" for c in value):
        raise ImportProblem(f"{where}: invalid or oversized cell.")
    if value.startswith(("=", "+", "-", "@")):
        raise ImportProblem(f"{where}: formula-like cells are not allowed; paste literal values.")
    return value


def table(rows, name):
    expected = SCHEMA[name]
    result = []
    number = 0
    for number, cells in enumerate(rows, 1):
        if number > MAX_ROWS + 1:
            raise ImportProblem(f"{name}: maximum {MAX_ROWS} data rows per sheet.")
        cells = list(cells)
        # openpyxl may pad rows to a formatted range. Only empty trailing cells
        # are ignored; CSV extra columns are never silently truncated.
        if len(cells) > len(expected) and any(v not in (None, "") for v in cells[len(expected):]):
            raise ImportProblem(f"{name} row {number}: unsupported columns.")
        if number > 1 and any(isinstance(v, int) and not isinstance(v, bool) and expected[i] != "version"
                              for i, v in enumerate(cells[:len(expected)])):
            raise ImportProblem(f"{name} row {number}: inventory keys and text must be text cells, not numbers.")
        values = [text_cell(v, f"{name} row {number}") for v in cells[:len(expected)]]
        values += [""] * (len(expected) - len(values))
        if number == 1:
            if tuple(values) != expected:
                raise ImportProblem(f"{name} row 1: headers must be exactly: {', '.join(expected)}.")
        elif any(values):
            result.append({"row": number, **dict(zip(expected, values))})
    if number:
        return result
    raise ImportProblem(f"{name}: missing header.")


def parse_upload(upload, batch):
    if upload.size > MAX_BYTES:
        raise ImportProblem("Maximum upload is 2 MiB.")
    data = upload.read(MAX_BYTES + 1)
    if len(data) > MAX_BYTES:
        raise ImportProblem("Maximum upload is 2 MiB.")
    suffix = upload.name.rsplit(".", 1)[-1].lower()
    if suffix == "csv":
        try:
            reader = csv.reader(io.StringIO(data.decode("utf-8-sig"), newline=""), strict=True)
            first = next(reader)
            if first != [VERSION, str(batch.pk), str(batch.space_id)]:
                raise ImportProblem("CSV row 1 must contain template version, exact installation UUID and space UUID.")
            plants = table(reader, "PlantUnits")
        except (UnicodeError, csv.Error, StopIteration) as exc:
            raise ImportProblem("CSV must be valid UTF-8 with the version/destination row and PlantUnits header.") from exc
        return validate_document({"Species": [], "CareGuides": [], "PlantUnits": plants})
    if suffix != "xlsx":
        raise ImportProblem("Use .xlsx or the inventory-only UTF-8 .csv template; legacy .xls and macros are unsupported.")
    try:
        with ZipFile(io.BytesIO(data)) as archive:
            infos = archive.infolist()
            if len(infos) > 80 or len({i.filename for i in infos}) != len(infos):
                raise ImportProblem("Workbook has too many or duplicate ZIP entries.")
            if sum(i.file_size for i in infos) > MAX_EXPANDED:
                raise ImportProblem("Expanded workbook exceeds 16 MiB.")
            for info in infos:
                name = info.filename.lower()
                if (info.flag_bits & 1 or ".." in name.split("/") or name.startswith(("/", "\\"))
                        or any(part in name for part in ("vbaproject", "externallinks/", "embeddings/", "connections.xml", "querytables/"))):
                    raise ImportProblem("Encrypted, macro, embedded or external-data workbooks are unsupported.")
                if name.endswith((".xml", ".rels")):
                    root = fromstring(archive.read(info))
                    for node in root.iter():
                        if node.tag.rsplit("}", 1)[-1] in ("f", "definedName"):
                            raise ImportProblem("Workbook formulas and defined names are unsupported; paste literal values.")
                        if node.attrib.get("TargetMode", "").lower() == "external":
                            raise ImportProblem("External workbook relationships are unsupported; use plain source URL text.")
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
        try:
            if set(book.sheetnames) != set(SCHEMA) or len(book.sheetnames) != 4:
                raise ImportProblem("Workbook must contain exactly Metadata, Species, CareGuides and PlantUnits sheets.")
            document = {}
            for name in SCHEMA:
                sheet = book[name]
                if sheet.max_row and sheet.max_row > MAX_ROWS + 1 or sheet.max_column and sheet.max_column > 20:
                    raise ImportProblem(f"{name}: sheet dimensions exceed the import limits.")
                sheet.reset_dimensions()  # Never trust producer-supplied dimensions.
                def values():
                    for row in sheet.iter_rows():
                        for cell in row:
                            if cell.data_type in ("f", "e", "b", "d"):
                                raise ImportProblem(f"{name} row {cell.row}: unsupported formula/error/boolean/date cell.")
                        yield [cell.value for cell in row]
                document[name] = table(values(), name)
        finally:
            book.close()
    except ImportProblem:
        raise
    except Exception as exc:
        # Parser exceptions must not echo arbitrary workbook contents to a log.
        raise ImportProblem("Invalid or unsupported XLSX workbook. Recreate it from the template.") from exc
    metadata = document.pop("Metadata")
    if len(metadata) != 3 or {r["field"]: r["value"] for r in metadata} != {
        "template_version": VERSION, "batch_id": str(batch.pk), "space_id": str(batch.space_id),
    }:
        raise ImportProblem("Metadata must match the template version and exact selected installation/space.")
    return validate_document(document)


def validate_document(document):
    if not document["PlantUnits"]:
        raise ImportProblem("PlantUnits: at least one real inventory row is required.")
    limits = {"source_key": 80, "label": 120, "species_label": 180, "placement_note": 180,
              "identity_reference": 500, "key": 100, "scientific_name": 180,
              "source_url": 500, "source_reference": 240}
    for name, rows in document.items():
        seen = set()
        for row in rows:
            where = f"{name} row {row['row']}"
            for key, value in row.items():
                if key in limits and len(value) > limits[key]:
                    raise ImportProblem(f"{where}: {key} exceeds {limits[key]} characters.")
            required = {"Species": ("key", "scientific_name"),
                        "CareGuides": ("species_key", "version", "source_url", "source_reference", "instructions"),
                        "PlantUnits": ("source_key", "label", "unit_type")}[name]
            if any(not row[key] for key in required):
                raise ImportProblem(f"{where}: required fields: {', '.join(required)}.")
            if name == "CareGuides":
                if not re.fullmatch(r"[1-9][0-9]{0,8}", row["version"]):
                    raise ImportProblem(f"{where}: version must be a positive integer.")
                row["version"] = int(row["version"])
                try:
                    if not row["source_url"].startswith("https://"):
                        raise ValidationError("HTTPS required")
                    URLValidator(schemes=["https"])(row["source_url"])
                except ValidationError as exc:
                    raise ImportProblem(f"{where}: a valid HTTPS source URL is required.") from exc
                key = (row["species_key"], row["version"])
            else:
                key = row["key"] if name == "Species" else row["source_key"]
            if key in seen:
                raise ImportProblem(f"{where}: duplicate key/version in this sheet.")
            seen.add(key)
            if name == "Species" and not re.fullmatch(r"[a-z0-9]+(?:-[a-z0-9]+)*", row["key"]):
                raise ImportProblem(f"{where}: species key must be a lowercase ASCII slug.")
            if name == "PlantUnits":
                if row["unit_type"] not in PlantUnit.UnitType.values:
                    raise ImportProblem(f"{where}: unit_type must be individual or group.")
                if bool(row["verified_species_key"]) != bool(row["identity_reference"]):
                    raise ImportProblem(f"{where}: verified_species_key and independently reviewed identity_reference are required together.")
    return document


def template_bytes(batch, csv_only=False):
    if csv_only:
        out = io.StringIO(newline="")
        writer = csv.writer(out)
        writer.writerow([VERSION, str(batch.pk), str(batch.space_id)])
        writer.writerow(SCHEMA["PlantUnits"])
        writer.writerow(["qa-unapproved-001", "FICTIONAL QA — NOT APPROVED", "individual", "Unknown", "", "", "QA test placement"])
        return out.getvalue().encode("utf-8-sig")
    book = Workbook()
    book.remove(book.active)
    examples = {
        "Metadata": [("template_version", VERSION), ("batch_id", str(batch.pk)), ("space_id", str(batch.space_id))],
        "Species": [("qa-fictional-species", "FICTIONAL QA SPECIES — NOT APPROVED")],
        "CareGuides": [("qa-fictional-species", 1, "https://example.test/qa-unapproved", "FICTIONAL QA ONLY — replace with an authorized source", "NOT BOTANICAL ADVICE. QA import example only.", "Not approved by Plantupia.")],
        "PlantUnits": [("qa-unapproved-001", "FICTIONAL QA — NOT APPROVED", "individual", "Unknown", "", "", "QA test placement")],
    }
    for name, headers in SCHEMA.items():
        sheet = book.create_sheet(name)
        sheet.append(headers)
        for row in examples[name]:
            sheet.append(row)
        sheet.freeze_panes = "A2"
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = 28
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def content_digest(document):
    # Row order/formatting/CSV vs XLSX do not change logical identity.
    return digest({k: sorted(({a: b for a, b in r.items() if a != "row"} for r in v),
                            key=lambda r: json.dumps(r, sort_keys=True)) for k, v in document.items()})


def plan_import(batch, document):
    """Read only. Existing taxonomy/version content can only match or conflict."""
    entries, errors, warnings = [], [], []
    incoming = {r["key"]: r for r in document["Species"]}
    keys = set(incoming) | {r["species_key"] for r in document["CareGuides"]} | {
        r["verified_species_key"] for r in document["PlantUnits"] if r["verified_species_key"]}
    species = {s.key: s for s in BotanicalSpecies.objects.filter(key__in=keys)}
    guides = {(g.species.key, g.version): g for g in CareGuideVersion.objects.filter(species__key__in=keys).select_related("species")}
    manifests = {r.source_key: r for r in batch.manifest_rows.select_related("plant__verified_species")}
    if any(r.plant.space_id != batch.space_id for r in manifests.values()):
        errors.append("Existing manifest does not match the exact destination space.")
    for name, rows in document.items():
        for row in rows:
            action, before, after = "create", {}, {k: v for k, v in row.items() if k != "row"}
            reason = ""
            if name == "Species":
                existing = species.get(row["key"])
                if existing:
                    before = {"key": existing.key, "scientific_name": existing.scientific_name}
                    action = "skip" if before == after else "conflict"
                    reason = "Existing curated identity must not be renamed by import."
            elif name == "CareGuides":
                existing = guides.get((row["species_key"], row["version"]))
                if row["species_key"] not in species and row["species_key"] not in incoming:
                    action, reason = "conflict", "Unknown species key. Include it in Species or curate it first."
                elif existing:
                    before = {k: getattr(existing, k) for k in SCHEMA[name] if k != "species_key"}
                    before["species_key"] = existing.species.key
                    action = "skip" if before == after else "conflict"
                    reason = "Existing versions are never overwritten; supply a new version."
                    warnings.append(f"CareGuides row {row['row']}: existing version is {'published' if existing.published else 'unpublished'}; publication state is preserved.")
                else:
                    reason = "Unreviewed draft; publication requires a separate staff action."
            else:
                key = row["verified_species_key"]
                if key and key not in species and key not in incoming:
                    action, reason = "conflict", "Unknown verified species key."
                existing = manifests.get(row["source_key"])
                if existing:
                    plant = existing.plant
                    before = {k: getattr(plant, k) for k in ("label", "unit_type", "species_label", "placement_note", "identity_reference")}
                    before.update(source_key=existing.source_key, verified_species_key=plant.verified_species.key if plant.verified_species_id else "")
                    if action != "conflict":
                        action = "skip" if before == after else "update"
                    if not plant.is_active:
                        action, reason = "conflict", "Archived units cannot be repurposed; replacements need a new source key and UUID."
                if key:
                    warnings.append(f"PlantUnits row {row['row']}: staff must independently review identity evidence and confirm verification.")
                else:
                    warnings.append(f"PlantUnits row {row['row']}: identity remains unverified; a free-text species is not proof.")
            entries.append(dict(sheet=name, row=row["row"], action=action, before=before, after=after, reason=reason))
            if action == "conflict":
                errors.append(f"{name} row {row['row']}: {reason}")
    # Include entire destination inventory plus referenced taxonomy/review state;
    # changes made after preview require another preview, never silent approval.
    state = {
        "space": str(batch.space_id), "owner": batch.space.owner_id, "status": batch.status,
        "manifest": list(batch.manifest_rows.order_by("pk").values("pk", "plant_id", "source_key")),
        "plants": list(PlantUnit.objects.filter(space_id=batch.space_id).order_by("pk").values()),
        "species": list(BotanicalSpecies.objects.filter(key__in=keys).order_by("key").values()),
        "guides": list(CareGuideVersion.objects.filter(species__key__in=keys).order_by("pk").values()),
    }
    receipt = InventoryImportReceipt.objects.filter(batch=batch, digest=content_digest(document)).first()
    if receipt:
        errors = []
        for entry in entries:
            entry.update(action="skip", reason="This content was already applied; existing records will not be restored or overwritten.")
    return dict(entries=entries, errors=errors, warnings=warnings, fingerprint=digest(state),
                summary=dict(Counter(e["action"] for e in entries)), replay=bool(receipt))


def make_preview(user, batch, document):
    authorize(user, batch)
    plan = plan_import(batch, document)
    token = None if plan["errors"] else signing.dumps({
        "actor": user.pk, "batch": str(batch.pk), "space": str(batch.space_id),
        "document": document, "fingerprint": plan["fingerprint"],
    }, salt=SALT, compress=True)
    return plan, token


def apply_preview(user, token, *, confirmed=False, identities_confirmed=False):
    require_permissions(user)
    if not confirmed:
        raise ImportProblem("Explicit import approval is required.")
    try:
        payload = signing.loads(token, salt=SALT, max_age=1800)
    except signing.BadSignature as exc:
        raise ImportProblem("Preview is invalid or expired. Upload again for a fresh preview.") from exc
    if payload["actor"] != user.pk:
        raise PermissionDenied
    document = payload["document"]
    if any(r["verified_species_key"] for r in document["PlantUnits"]) and not identities_confirmed:
        raise ImportProblem("Explicit confirmation of independently reviewed identity evidence is required.")
    try:
        with transaction.atomic():
            # Same parent-first lock order as inventory review and handover.
            batch = InstallationBatch.objects.select_for_update().get(pk=payload["batch"])
            batch.space = Space.objects.select_for_update().get(pk=batch.space_id)
            authorize(user, batch)
            if str(batch.space_id) != payload["space"]:
                raise ImportProblem("Destination changed. Upload for a fresh preview.")
            receipt = InventoryImportReceipt.objects.filter(batch=batch, digest=content_digest(document)).first()
            if receipt:
                return receipt, False
            list(PlantUnit.objects.select_for_update().filter(space_id=batch.space_id).order_by("pk"))
            # Species locks serialize guide publication/import across batches.
            keys = sorted(set(r["key"] for r in document["Species"]) |
                          set(r["species_key"] for r in document["CareGuides"]) |
                          set(r["verified_species_key"] for r in document["PlantUnits"] if r["verified_species_key"]))
            list(BotanicalSpecies.objects.select_for_update().filter(key__in=keys).order_by("key"))
            list(CareGuideVersion.objects.select_for_update().filter(species__key__in=keys).order_by("pk"))
            plan = plan_import(batch, document)
            if plan["errors"] or plan["fingerprint"] != payload["fingerprint"]:
                raise ImportProblem("Inventory or guide review changed since preview. Upload again; no rows were applied.")
            for row in sorted(document["Species"], key=lambda r: r["key"]):
                if not BotanicalSpecies.objects.filter(key=row["key"]).exists():
                    obj = BotanicalSpecies(key=row["key"], scientific_name=row["scientific_name"])
                    obj.full_clean()
                    obj.save()
            for row in sorted(document["CareGuides"], key=lambda r: (r["species_key"], r["version"])):
                species = BotanicalSpecies.objects.get(key=row["species_key"])
                if not CareGuideVersion.objects.filter(species=species, version=row["version"]).exists():
                    CareGuideVersion.objects.create(species=species, **{
                        key: row[key] for key in SCHEMA["CareGuides"] if key != "species_key"})
            for row in document["PlantUnits"]:
                entry = next(e for e in plan["entries"] if e["sheet"] == "PlantUnits" and e["row"] == row["row"])
                if entry["action"] == "skip":
                    continue
                manifest = batch.manifest_rows.filter(source_key=row["source_key"]).select_related("plant").first()
                plant = manifest.plant if manifest else PlantUnit(space=batch.space)
                for key in ("label", "unit_type", "species_label", "placement_note"):
                    setattr(plant, key, row[key])
                new_species = BotanicalSpecies.objects.get(key=row["verified_species_key"]) if row["verified_species_key"] else None
                identity_changed = plant.verified_species_id != (new_species.pk if new_species else None) or plant.identity_reference != row["identity_reference"]
                plant.verified_species = new_species
                plant.identity_reference = row["identity_reference"]
                if identity_changed or new_species and not plant.identity_verified_by_id:
                    plant.identity_verified_by = user if new_species else None
                    plant.identity_verified_at = timezone.now() if new_species else None
                plant.full_clean()
                plant.save()
                if not manifest:
                    manifest = InstallationManifestRow(batch=batch, plant=plant, source_key=row["source_key"])
                    manifest.full_clean()
                    manifest.save()
            receipt = InventoryImportReceipt.objects.create(batch=batch, digest=content_digest(document), approved_by=user, summary=plan["summary"])
            return receipt, True
    except (IntegrityError, ValidationError, InstallationBatch.DoesNotExist, Space.DoesNotExist) as exc:
        raise ImportProblem("Import conflicted with current records. No rows applied; upload for a fresh preview.") from exc
