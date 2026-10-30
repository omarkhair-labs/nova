"""Opt-in QA-only ZIP catalogue import. No real commerce or production fallback."""
from __future__ import annotations

import os
import re
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import PurePosixPath
from zipfile import BadZipFile, ZipFile

from django import forms
from django.contrib import messages
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import transaction
from django.http import Http404
from django.shortcuts import redirect
from django.utils.text import slugify
from openpyxl import load_workbook

from . import company_services as company
from .company_demo import demo_portal
from .company_views import page
from .media import sanitize_public_plant_photo
from .models import CatalogProduct

MAX_ZIP_BYTES = 20 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 40 * 1024 * 1024
MAX_IMAGE_BYTES = 2 * 1024 * 1024
REQUIRED = ("sku", "name", "kind", "description", "proposed_price", "currency",
            "qa_stock", "active", "image", "demo_only")


class CatalogPackError(ValueError):
    pass


class CatalogZipForm(forms.Form):
    upload = forms.FileField(help_text="ZIP with products.xlsx and images/*.webp, 20 MiB maximum.")
    action = forms.ChoiceField(choices=[("preview", "Validate only (no changes)"),
                                         ("import", "Confirm import into isolated QA")])

    def clean_upload(self):
        file = self.cleaned_data["upload"]
        if file.size > MAX_ZIP_BYTES:
            raise forms.ValidationError("ZIP exceeds the 20 MiB limit.")
        if not file.name.lower().endswith(".zip"):
            raise forms.ValidationError("Upload a ZIP package.")
        return file


@dataclass
class ProposedProduct:
    sku: str
    name: str
    kind: str
    description: str
    price: Decimal
    qa_stock: int
    active: bool
    photo: SimpleUploadedFile


def _flag(value, field: str) -> bool:
    if value is True or str(value).strip().lower() in {"true", "1", "yes"}:
        return True
    if value is False or str(value).strip().lower() in {"false", "0", "no"}:
        return False
    raise CatalogPackError(f"Invalid {field}; use TRUE or FALSE.")


def parse_pack(upload) -> list[ProposedProduct]:
    if upload.size > MAX_ZIP_BYTES:
        raise CatalogPackError("ZIP exceeds 20 MiB.")
    try:
        with ZipFile(upload) as archive:
            infos = archive.infolist()
            if len(infos) > 90 or sum(i.file_size for i in infos) > MAX_UNCOMPRESSED_BYTES:
                raise CatalogPackError("ZIP contains too much decompressed data.")
            paths = [i.filename for i in infos if not i.is_dir()]
            if len(paths) != len(set(paths)):
                raise CatalogPackError("ZIP has repeated file names.")
            if any(not p or p.startswith("/") or "\\" in p
                   or ".." in PurePosixPath(p).parts for p in paths):
                raise CatalogPackError("Invalid ZIP paths.")
            if "products.xlsx" not in paths:
                raise CatalogPackError("ZIP must include products.xlsx at its root.")
            workbook = load_workbook(BytesIO(archive.read("products.xlsx")), read_only=True, data_only=True)
            try:
                if "Products" not in workbook.sheetnames:
                    raise CatalogPackError("Workbook must contain a Products sheet.")
                worksheet = workbook["Products"]
                # Streaming XLSX files may omit row dimension metadata (max_row=None).
                # The loop below independently enforces the 100-product limit.
                if worksheet.max_row is not None and worksheet.max_row > 101:
                    raise CatalogPackError("Maximum 100 catalogue rows per import.")
                values = worksheet.iter_rows(values_only=True)
                header = tuple(str(v).strip() if v is not None else "" for v in next(values, ()))
                if header[:len(REQUIRED)] != REQUIRED:
                    raise CatalogPackError("Column headers do not match the Plantupia demo template.")
                parsed = []
                skus, names = set(), set()
                for number, fields in enumerate(values, 2):
                    if not any(v is not None and str(v).strip() for v in fields):
                        continue
                    if len(parsed) >= 100:
                        raise CatalogPackError("Maximum 100 products per import.")
                    vals = dict(zip(REQUIRED, fields))
                    sku = str(vals["sku"] or "").strip()
                    name = str(vals["name"] or "").strip()
                    kind = str(vals["kind"] or "").strip()
                    description = str(vals["description"] or "").strip()
                    image_path = str(vals["image"] or "").strip()
                    if not re.fullmatch(r"[A-Z0-9-]{3,32}", sku):
                        raise CatalogPackError(f"Row {number}: invalid SKU.")
                    if sku in skus or name in names:
                        raise CatalogPackError(f"Row {number}: duplicate SKU or product name.")
                    if not name.startswith("FICTIONAL QA") or len(name) > 120 or not description or len(description) > 1500:
                        raise CatalogPackError(f"Row {number}: use a fictional QA name and a description.")
                    if kind not in {"plant", "planter", "accessory"} or str(vals["currency"]) != "EGP":
                        raise CatalogPackError(f"Row {number}: invalid kind or currency.")
                    if not _flag(vals["demo_only"], "demo_only"):
                        raise CatalogPackError(f"Row {number}: commercial products are not accepted.")
                    active = _flag(vals["active"], "active")
                    try:
                        price = Decimal(str(vals["proposed_price"]))
                        if not price.is_finite() or price < 0 or price > Decimal("999999.99") or price.as_tuple().exponent < -2:
                            raise ValueError
                        qty = int(vals["qa_stock"])
                        if str(qty) != str(vals["qa_stock"]).replace(".0", "") or not 0 <= qty <= 10000:
                            raise ValueError
                    except (InvalidOperation, ValueError, TypeError):
                        raise CatalogPackError(f"Row {number}: invalid provisional price or QA stock.")
                    if image_path not in paths or not image_path.startswith("images/") or image_path.count("/") != 1:
                        raise CatalogPackError(f"Row {number}: image does not exist inside images/.")
                    if not image_path.lower().endswith((".webp", ".png", ".jpg", ".jpeg")):
                        raise CatalogPackError(f"Row {number}: invalid image extension.")
                    img_info = archive.getinfo(image_path)
                    if img_info.file_size > MAX_IMAGE_BYTES:
                        raise CatalogPackError(f"Row {number}: image exceeds 2 MiB.")
                    photo = SimpleUploadedFile(image_path.rsplit("/", 1)[-1], archive.read(image_path))
                    try:
                        photo = sanitize_public_plant_photo(photo)
                    except ValueError as exc:
                        raise CatalogPackError(f"Row {number}: image cannot be normalized ({exc}).") from exc
                    if photo.size > MAX_IMAGE_BYTES:
                        raise CatalogPackError(f"Row {number}: normalized image exceeds 2 MiB.")
                    parsed.append(ProposedProduct(sku, name, kind, description, price, qty, active, photo))
                    skus.add(sku); names.add(name)
                if not parsed:
                    raise CatalogPackError("No products in the package.")
                return parsed
            finally:
                workbook.close()
    except (BadZipFile, OSError, RuntimeError) as exc:
        raise CatalogPackError("Cannot read ZIP package.") from exc


def import_products(user, entries: list[ProposedProduct]) -> int:
    """DB atomic; remove any created image objects on error, never modify existing items."""
    stored = []
    try:
        with transaction.atomic():
            existing = set(CatalogProduct.objects.select_for_update().filter(
                name__in=[p.name for p in entries]).values_list("name", flat=True))
            existing_skus = set(CatalogProduct.objects.select_for_update().filter(
                sku__in=[p.sku for p in entries]).values_list("sku", flat=True))
            if existing_skus:
                raise CatalogPackError("Existing SKU values would be duplicated: " +
                    ", ".join(sorted(existing_skus))[:300])
            if existing:
                raise CatalogPackError("Already exists: " + ", ".join(sorted(existing))[:300])
            for entry in entries:
                obj = CatalogProduct(name=entry.name, sku=entry.sku, kind=entry.kind,
                    description=entry.description, proposed_price=entry.price,
                    currency="EGP", qa_stock=entry.qa_stock, active=entry.active)
                obj.full_clean(exclude=["image"])
                safe_name = (slugify(entry.sku) or "qa") + ".jpg"
                obj.image.save(safe_name, entry.photo, save=False)
                stored.append((obj.image.storage, obj.image.name))
                obj.save()
                company.audit(user, obj, f"QA ZIP import: {entry.sku}; concept photo and provisional price only")
        return len(entries)
    except Exception:
        for storage, name in stored:
            try:
                storage.delete(name)
            except Exception:
                pass
        raise


@demo_portal("plants.add_catalogproduct")
def catalog_zip_import(request):
    # Never auto-activate this feature with the existing CEO demo switch.
    if os.getenv("PLANTUPIA_QA_CATALOG_ZIP_IMPORT_ENABLED", "0") != "1":
        raise Http404
    form = CatalogZipForm(request.POST or None, request.FILES or None)
    preview, outcome = [], None
    if request.method == "POST" and form.is_valid():
        try:
            proposed = parse_pack(form.cleaned_data["upload"])
            existing = set(CatalogProduct.objects.filter(
                name__in=[p.name for p in proposed]).values_list("name", flat=True))
            if existing:
                raise CatalogPackError("Existing products would be duplicated: " + ", ".join(sorted(existing))[:300])
            preview = proposed
            if form.cleaned_data["action"] == "import":
                number = import_products(request.user, proposed)
                messages.success(request, f"Imported {number} QA concepts and photos. No payment or delivery is enabled.")
                return redirect("company:catalog")
            outcome = f"Validated {len(proposed)} products and all images. Nothing was saved."
        except (CatalogPackError, ValidationError) as exc:
            form.add_error("upload", str(exc))
    return page(request, "catalog_zip_import", {
        "title": "Import fictional QA catalogue", "section": "catalog",
        "form": form, "preview": preview, "outcome": outcome,
    }, status=400 if form.errors else 200)
